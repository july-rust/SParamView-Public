from pathlib import Path
import json,math,zipfile
import numpy as np
from scipy.special import i0
from scipy.integrate import cumulative_trapezoid
import openpyxl
root=Path('real-validation');fileinfo=json.loads((root/'real_file_crosscheck.json').read_text())['files'];networks={}
for item in fileinfo:
 p=Path(item['probe'])/'matrix.bin';n,k=item['ports'],item['points']
 f=np.memmap(p,dtype='<f8',mode='r',offset=16,shape=(k,))
 s=np.memmap(p,dtype='<c16',mode='r',offset=16+k*8,shape=(n,n,k))
 networks[item['name']]=(f,s)

def endpoint(c,far=False):
 p,n=c['farP' if far else 'nearP'],c['farN' if far else 'nearN']
 return [(p,1)] if n<0 else [(p,1/np.sqrt(2)),(n,-1/np.sqrt(2))]
def extract(net,c,metric,ag=None):
 f,s=net;response=endpoint(c,metric in ('Insertion Loss','FEXT'));stimulus=endpoint(ag if ag else c)
 v=np.zeros(len(f),complex)
 for i,a in response:
  for j,b in stimulus:v+=a*b*s[i,j]
 return np.asarray(f),v

def crop(f,v,start,stop):
 lo=max(start,float(f[0]));hi=min(stop,float(f[-1]));x=np.r_[lo,f[(f>lo)&(f<hi)],hi] if hi>lo else np.array([lo])
 y=np.interp(x,f,v.real)+1j*np.interp(x,f,v.imag)
 return x,y

def tdr(f,v,settings,ref):
 f,v=crop(f,v,f[0],settings['stopHz']);nominal=(f[-1]-f[0])/(len(f)-1);bins=math.ceil(f[-1]/nominal);n=16
 while n//2<bins:n*=2
 df=f[-1]/(n//2);grid=np.arange(n//2+1)*df
 dc=v[0].real if f[0]==0 else v[0].real+(v[0].real-v[1].real)*f[0]/(f[1]-f[0])
 gf=np.r_[0,f] if f[0]>0 else f;gs=np.r_[dc,v] if f[0]>0 else v.copy();gs[0]=dc
 spec=np.interp(grid,gf,gs.real)+1j*np.interp(grid,gf,gs.imag)
 approx=np.interp(f,grid,spec.real)+1j*np.interp(f,grid,spec.imag);err=float(np.max(abs(approx-v)))
 if err>.05:return None,None,err
 u=np.arange(n//2+1)/(n//2);spec*=i0(settings['kaiserBeta']*np.sqrt(np.maximum(0,1-u*u)))/i0(settings['kaiserBeta'])
 spec[0]=spec[0].real;spec[-1]=spec[-1].real
 impulse=np.fft.fftshift(np.fft.irfft(spec,n=n));rho=cumulative_trapezoid(impulse,initial=0)
 times=np.fft.fftshift(np.fft.fftfreq(n,d=df));z=ref*(1+rho)/(1-rho)
 z[(abs(1-rho)<1e-8)|(abs(rho)>1.0001)]=np.nan
 mask=(times>=max(0,settings['tdrStart']))&(times<=min(.45/df,settings['tdrStop']))
 return times[mask],z[mask],err

checks=[];summaries=[]
for case in json.loads((root/'cases.json').read_text()):
 key=case['key'];p=json.loads(Path(case['project']).read_text());g=root/'gui'/key;j=json.loads((g/'Validation.json').read_text());assert j['results']==case['expected_results']
 revs={r['name']:r for r in p['revisions']};channels={r['name']:{c['name']:c for c in r['channels']} for r in p['revisions']};cfg=p['settings'];data={}
 for row in j['analysis']:
  directions={'Return Loss':'Near -> Near','Insertion Loss':'Near -> Far','NEXT':'Aggressor Near -> Victim Near','FEXT':'Aggressor Near -> Victim Far','TDR':'Near -> Near'}
  assert row['direction']==directions[row['metric']]
  rev,cn,m=row['revision'],row['channel'],row['metric'];c=channels[rev][cn];ag=channels[rev].get(row['aggressor']);net=networks[Path(revs[rev]['source']).name];f,v=extract(net,c,m,ag)
  if m=='TDR':
   x,y,err=tdr(f,v,cfg,100 if c['nearN']>=0 else 50)
   if x is None:
    assert row['worst'] is None and row['quality']=='UNSUITABLE' and 'grid loses' in row['note'],row
    checks.append(dict(case=key,channel=cn,revision=rev,metric=m,result='Correctly unavailable',resampling_error=err));continue
   target=cfg['targetOhm'] or (100 if c['nearN']>=0 else 50);i=int(np.nanargmax(abs(y-target)));tol=1e-6
  else:
   x,v=crop(f,v,cfg['startHz'],cfg['stopHz']);y=20*np.log10(abs(v));i=int(np.argmin(y) if m=='Insertion Loss' else np.argmax(y));tol=1e-7
   data[(rev,cn,m,row['aggressor'])]=(x,v)
  delta=abs(float(y[i])-row['worst']);assert delta<tol,(key,cn,m,delta,row['worst'],y[i])
  assert abs(float(np.nanmin(y))-row['minimum'])<tol
  assert abs(float(np.nanmax(y))-row['maximum'])<tol
  assert row['status']=='N/A' and row['margin'] is None
  checks.append(dict(case=key,channel=cn,revision=rev,metric=m,max_stat_error=float(delta)))
 for metric,rows in j.get('comparisons',{}).items():
  for row in rows:
   label=row['Channel / Aggressor'].split(' ← ');cn=label[0];ag=label[1] if len(label)>1 else ''
   af,av=data[(row['Baseline'],cn,metric,ag)];bf,bv=data[(row['Candidate'],cn,metric,ag)]
   grid=np.union1d(af,bf);aa=np.interp(grid,af,av.real)+1j*np.interp(grid,af,av.imag);bb=np.interp(grid,bf,bv.real)+1j*np.interp(grid,bf,bv.imag)
   ay=20*np.log10(abs(av));by=20*np.log10(abs(bv));worst=lambda y:float(np.min(y) if metric=='Insertion Loss' else np.max(y))
   raw=worst(by)-worst(ay);sign=1 if metric=='Insertion Loss' else -1
   pointwise=sign*(20*np.log10(abs(bb))-20*np.log10(abs(aa)))
   for name,wanted in [('Base Worst',worst(ay)),('Candidate Worst',worst(by)),('Raw Δ dB',raw),('Improvement dB',sign*raw),('Min pointwise improvement',float(np.min(pointwise)))]:
    assert abs(float(row[name])-wanted)<.000501,(key,metric,cn,name,row[name],wanted)
   checks.append(dict(case=key,metric=metric,comparison=cn+' <- '+ag,result='PASS at displayed 0.001 dB precision'))
 with zipfile.ZipFile(g/'Analysis.xlsx') as z:assert z.testzip() is None;images=sum(n.startswith('xl/media/') for n in z.namelist())
 wb=openpyxl.load_workbook(g/'Analysis.xlsx',read_only=True,data_only=False);detail_count=sum(sum(1 for _ in wb[name].iter_rows())-1 for name in ['Return_Loss','Insertion_Loss','NEXT','FEXT','TDR']);assert detail_count==j['results']
 assert 'Raw_1' not in wb.sheetnames
 summaries.append(dict(case=key,results=j['results'],xlsx_detail_rows=detail_count,xlsx_graph_images=images,gui_seconds=j['elapsed_seconds'],quality_counts={q:sum(r['metric']=='TDR' and r['quality']==q for r in j['analysis']) for q in ['GOOD','LIMITED','UNSUITABLE']}))
 print(summaries[-1],flush=True)
report=dict(analysis_checks=len(checks),cases=summaries,checks=checks)
(root/'analysis_crosscheck.json').write_text(json.dumps(report,indent=2))
print('PASS',len(checks),flush=True)
